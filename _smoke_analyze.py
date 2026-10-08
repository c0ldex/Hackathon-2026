import io
from openpyxl import Workbook

from analyzer import analyze_workbook
from narrative import enrich_with_narrative
from report import render_executive_report

wb = Workbook()
ws = wb.active
ws.title = "CVE"
ws.append(
    [
        "vulnerability.id",
        "vulnerability.severity",
        "vulnerability.score.base",
        "agent.name",
        "package.name",
        "vulnerability.description",
    ]
)
ws.append(["CVE-2024-0001", "Critical", 9.8, "lab-host-1", "openssl", "synthetic lab finding"])

header = "timestamp,rule.level,rule.description,agent.name,rule.id"
ws2 = wb.create_sheet("Windows")
ws2.append([header])
ws2.append(['Jan 2, 2026 @ 10:00:00,12,"stop auditd",lab-host-1,100'])
ws2.append(['Jan 2, 2026 @ 10:01:00,10,"privileged command root;",lab-host-1,101'])

ws3 = wb.create_sheet("Fortigate")
ws3.append([header])
ws3.append(['Jan 2, 2026 @ 10:02:00,8,"connection failed",lab-host-1,200'])

buf = io.BytesIO()
wb.save(buf)
result = enrich_with_narrative(analyze_workbook(buf.getvalue(), filename="planilha.xlsx"))
html = render_executive_report(result)
print("score", result["finalRiskScore"], result["riskLevel"])
print("sources", result["sources"])
print("recs", len(result["recommendations"]))
print("evidence", [item["id"] for item in result["evidenceTrail"]])
print("checks", result["consistency"]["allChecksPassed"])
print("narrative", result.get("narrativeSource"))
print("html_len", len(html))
print("r1", result["correlation"]["hostsWithCriticalCveAndHighEvent"][:1])
